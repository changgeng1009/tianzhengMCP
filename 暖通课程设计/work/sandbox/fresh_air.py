import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment

wb = openpyxl.Workbook()
ws = wb.active
ws.title = '新风量计算-部分'

# Parameters
H = 4.2  # 净高 m (待确认)
h_enthalpy = 28.8  # kJ/kg 南宁地区

headers = ['房间','面积m2','净高m','人员密度','人数','人均新风','人员法Q1','换气次数','换气法Q2','总送风量L','L*10%','最终新风量','新风负荷kW']
for c, h in enumerate(headers, 1):
    ws.cell(row=1, column=c, value=h).font = Font(bold=True)

# 中餐厅(左): 286.68 m2
A1 = 286.68
N1 = A1 * 0.80  # 229 people
Q1_1 = N1 * 25   # 5733.6
Q2_1 = 4 * A1 * H  # 4816.2
L1 = 12 * A1 * H   # 14448.7
Q3_1 = L1 * 0.10   # 1444.9
Q_final_1 = max(Q1_1, Q2_1, Q3_1)
QL_1 = 1.2 * Q_final_1 * h_enthalpy / 3600

# 中餐厅(右): 432.01 m2
A2 = 432.01
N2 = A2 * 0.80
Q1_2 = N2 * 25
Q2_2 = 4 * A2 * H
L2 = 12 * A2 * H
Q3_2 = L2 * 0.10
Q_final_2 = max(Q1_2, Q2_2, Q3_2)
QL_2 = 1.2 * Q_final_2 * h_enthalpy / 3600

rows = [
    ['中餐厅(左)', A1, H, 0.80, round(N1,0), 25, round(Q1_1,1), 4, round(Q2_2 if False else Q2_1,1), round(L1,1), round(Q3_1,1), round(Q_final_1,1), round(QL_1,2)],
    ['中餐厅(右)', A2, H, 0.80, round(N2,0), 25, round(Q1_2,1), 4, round(Q2_2,1), round(L2,1), round(Q3_2,1), round(Q_final_2,1), round(QL_2,2)],
]

# Fixed Q2_1
rows[0][8] = round(4 * A1 * H, 1)

for r, row in enumerate(rows, 2):
    for c, v in enumerate(row, 1):
        ws.cell(row=r, column=c, value=v).alignment = Alignment(horizontal='center')

for i in range(1, 14):
    ws.column_dimensions[get_column_letter(i)].width = 14

# Add notes
ws.cell(row=5, column=1, value='参数说明：').font = Font(bold=True)
ws.cell(row=6, column=1, value='净高=4.2m (指导书示例值, 待DWG确认)')
ws.cell(row=7, column=1, value='人员密度=0.80人/m2 (指导书推荐)')
ws.cell(row=8, column=1, value='人均新风=25m3/h人 (JGJ64-2017)')
ws.cell(row=9, column=1, value='新风焓差=28.8kJ/kg (南宁地区)')
ws.cell(row=10, column=1, value='空调换气=12次/h, 新风换气=4次/h')

wb.save(r'D:\CodexWork\天正\暖通课程设计\work\partial_fresh_air_calculation.xlsx')
print('Created partial_fresh_air_calculation.xlsx')
