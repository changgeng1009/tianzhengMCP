import openpyxl
from openpyxl.utils import get_column_letter

# Update room_area_extraction.xlsx
wb = openpyxl.load_workbook(r'D:\CodexWork\天正\暖通课程设计\work\room_area_extraction.xlsx')
ws = wb['房间面积提取']

# Update known areas
ws.cell(row=2, column=6, value=286.68)  # 中餐厅左
ws.cell(row=2, column=11, value='用户提供')
ws.cell(row=3, column=6, value=432.01)  # 中餐厅右
ws.cell(row=3, column=11, value='用户提供')

# Add height estimate (4.2m from course guide example, pending DWG confirmation)
ws.cell(row=2, column=7, value=4.2)
ws.cell(row=3, column=7, value=4.2)

wb.save(r'D:\CodexWork\天正\暖通课程设计\work\room_area_extraction.xlsx')
print('Updated room_area_extraction.xlsx')

# Now create partial load calculation
wb2 = openpyxl.Workbook()
ws2 = wb2.active
ws2.title = '冷负荷计算-部分'

headers = ['房间', '面积m2', '冷负荷指标W/m2', '房间冷负荷kW', '选型冷量kW', '空调形式', '备注']
for c, h in enumerate(headers, 1):
    ws2.cell(row=1, column=c, value=h).font = openpyxl.styles.Font(bold=True)

rooms = [
    ['中餐厅(左)', 286.68, 430, 286.68*430/1000, 286.68*430/1000*1.1, '全空气系统', '面积已确认'],
    ['中餐厅(右)', 432.01, 430, 432.01*430/1000, 432.01*430/1000*1.1, '全空气系统', '面积已确认'],
]
for r, row in enumerate(rooms, 2):
    for c, v in enumerate(row, 1):
        cell = ws2.cell(row=r, column=c, value=round(v, 2) if isinstance(v, float) else v)
        cell.alignment = openpyxl.styles.Alignment(horizontal='center')

for i in range(1, 8):
    ws2.column_dimensions[get_column_letter(i)].width = 16

wb2.save(r'D:\CodexWork\天正\暖通课程设计\work\partial_load_calculation.xlsx')
print('Created partial_load_calculation.xlsx')
