# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook
src=Path('output/暖通课程设计计算表-GPT重填版.xlsx')
patched=Path('output/暖通课程设计计算表-GPT重填版-模板内容补丁.xlsx')
wb0=load_workbook(src, data_only=False)
wb=load_workbook(patched, data_only=False)
for s in wb0.sheetnames:
    f0={}; f={}
    for row in wb0[s].iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith('='): f0[c.coordinate]=c.value
    for row in wb[s].iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith('='): f[c.coordinate]=c.value
    lost=[coord for coord in f0 if coord not in f]
    if lost:
        print('SHEET',s)
        for coord in lost:
            print(coord, f0[coord], 'patched=', wb[s][coord].value)
wb0.close(); wb.close()
