# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook
base=Path.cwd()
calc=base/'output'/'暖通课程设计计算表-GPT重填版.xlsx'
template=base/'work'/'sandbox'/'计算表-原始模板备份.xlsx'
print(calc, calc.exists(), calc.stat().st_size if calc.exists() else None)
for file in [calc, template]:
    wb=load_workbook(file, read_only=False, data_only=False)
    print('\nFILE',file.name)
    for ws in wb.worksheets:
        print('\nSHEET',ws.title,ws.max_row,ws.max_column,'merged',len(ws.merged_cells.ranges))
        # print nonempty cells in first 30 rows and 25 cols compact
        for r in range(1, min(ws.max_row,30)+1):
            vals=[]
            for c in range(1, min(ws.max_column,25)+1):
                v=ws.cell(r,c).value
                if v is not None:
                    vals.append(f'{ws.cell(r,c).coordinate}={str(v)[:60]}')
            if vals:
                print(' | '.join(vals))
    wb.close()
