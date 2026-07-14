# -*- coding: utf-8 -*-
from pathlib import Path
from shutil import copy2
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from datetime import datetime
base=Path.cwd()
src=base/'output'/'暖通课程设计计算表-GPT重填版-模板内容补丁.xlsx'
out=base/'output'/'暖通课程设计计算表-GPT重填版-清洁补丁.xlsx'
copy2(src,out)
terms=['预留','需人工确认','待人工确认','按平面布置','按平面布置估算','待图纸布置后复核','待施工图复核','人工确认']
wb=load_workbook(out, data_only=False)
removed=[]
# clean all disallowed remark strings from normal cells
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            v=cell.value
            if isinstance(v,str) and any(t in v for t in terms):
                removed.append({'sheet':ws.title,'cell':cell.coordinate,'value':v})
                cell.value=None
# explicitly clean water length columns C/K in data regions if text remains, preserving headers and unit rows
ws=wb['水系统']
for col in [3,11]:
    for r in range(23, ws.max_row+1):
        cell=ws.cell(r,col)
        if isinstance(cell.value,str):
            if r not in [33,34,45,46] and cell.value not in ['长度','m']:
                removed.append({'sheet':'水系统','cell':cell.coordinate,'value':cell.value})
                cell.value=None
# Fill indoor air design parameter table in existing rows B18:H21
ws=wb['室内外设计参数']
indoor_rows=[
    (18,'中餐厅',25,55,18,45,25,4),
    (19,'豪华包厢',25,55,18,45,25,4),
    (20,'普通包厢',25,55,18,45,25,4),
    (21,'标准客房',25,55,18,45,30,None),
]
for r,name,summer_t,summer_h,winter_t,winter_h,fresh,ach in indoor_rows:
    vals=[name,summer_t,summer_h,winter_t,winter_h,fresh,ach]
    for c,v in enumerate(vals, start=2):
        ws.cell(r,c).value=v
# Preserve calculation behavior on open
try:
    wb.calculation.fullCalcOnLoad=True
    wb.calculation.forceFullCalc=True
    wb.calculation.calcMode='auto'
except Exception:
    pass
wb.save(out)
wb.close()
# Markdown remarks
lines=['# Excel备注与人工确认项','',f'- 生成时间：{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}',f'- 清洁补丁文件：`{out}`','- 本文件集中记录从 Excel 表格单元格中移出的说明性文字。Excel 数值表内不再写入“预留”“需人工确认”“按平面布置”等备注。','','## 已从 Excel 单元格清理的备注','', '| 工作表 | 单元格 | 原内容 |','|---|---|---|']
for item in removed:
    val=str(item['value']).replace('|','/')
    lines.append(f"| {item['sheet']} | {item['cell']} | {val} |")
lines += ['','## 集中说明','', '1. 水系统表的“长度 m”列在当前课程设计模板公式中不作为关键阻塞项；清洁补丁中已清空文字备注，后续如老师要求，可按实际平面布置补入数字。', '2. 4F 层高如老师要求严格核对，可继续结合图纸或教师意见复核。', '3. 客房冷负荷指标 120 W/m² 为课程设计取值；如老师要求考虑朝向、窗墙比或西晒，可再调整。', '4. 设备最终型号命名按老师样本复核；当前 Excel 已采用本地资料提取的 FC、KD、ZK 系列。', '5. 风口位置和管线布置属于后续 CAD 图纸阶段内容，不写入当前 Excel 数值表。']
(base/'output'/'Excel备注与人工确认项.md').write_text('\n'.join(lines), encoding='utf-8')
print('cleaned', len(removed), 'cells')
print(out)
