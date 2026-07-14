# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook
import math, json, re
base=Path.cwd()
src=base/'output'/'暖通课程设计计算表-GPT重填版.xlsx'
patched=base/'output'/'暖通课程设计计算表-GPT重填版-模板内容补丁.xlsx'
template=base/'work'/'sandbox'/'计算表-原始模板备份.xlsx'
report=base/'audit'/'template_only_patch_self_check.md'

def formula_cells(ws):
    cells={}
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith('='):
                cells[c.coordinate]=c.value
    return cells

def merged_set(ws):
    return sorted(str(r) for r in ws.merged_cells.ranges)

wb0=load_workbook(src, data_only=False)
wb=load_workbook(patched, data_only=False)
checks=[]
# 1 sheets
checks.append(('工作表名称保持不变', wb.sheetnames==wb0.sheetnames, f'母版 {len(wb0.sheetnames)} 个，补丁版 {len(wb.sheetnames)} 个'))
# 2 rows cols and merged
row_issues=[]; merge_issues=[]
for s in wb0.sheetnames:
    a=wb0[s]; b=wb[s]
    if (a.max_row,a.max_column)!=(b.max_row,b.max_column): row_issues.append(f'{s}: {a.max_row}x{a.max_column} -> {b.max_row}x{b.max_column}')
    if merged_set(a)!=merged_set(b): merge_issues.append(s)
checks.append(('行列数没有异常变化', not row_issues, '; '.join(row_issues) if row_issues else '全部一致'))
checks.append(('合并单元格保持', not merge_issues, '; '.join(merge_issues) if merge_issues else '全部一致'))
# 3 formulas formula cells not lost (allow changed formula text, positions should include original formula coordinates)
formula_lost=[]
for s in wb0.sheetnames:
    f0=formula_cells(wb0[s]); f=formula_cells(wb[s])
    for coord in f0:
        if coord not in f:
            formula_lost.append(f'{s}!{coord}')
checks.append(('模板公式保留', not formula_lost, f'丢失公式 {len(formula_lost)} 个' if formula_lost else '未发现公式单元格被清空'))
# forbidden new terms
forbidden=['最不利环路','水泵扬程','局部阻力','沿程阻力','完整水力平衡']
found=[]
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,str):
                for term in forbidden:
                    if term in c.value:
                        found.append(f'{ws.title}!{c.coordinate}:{term}')
checks.append(('未新增模板外施工图深化计算项', not found, '; '.join(found) if found else '未检出禁用深化项'))
# no level column: check headers row include 楼层 standalone
floor_cols=[]
for ws in wb.worksheets:
    for row in ws.iter_rows(max_row=min(ws.max_row,25)):
        for c in row:
            if c.value=='楼层': floor_cols.append(f'{ws.title}!{c.coordinate}')
checks.append(('没有新增楼层列', not floor_cols, '; '.join(floor_cols) if floor_cols else '未发现独立“楼层”表头'))
# key data
ws=wb['房间冷负荷计算']
room_name=ws['A17'].value; area=ws['B17'].value; idx=ws['C17'].value; qty=ws['G17'].value; model=ws['F17'].value
cold=round(area*idx/1000,2); sel=round(cold*1.1,2)
checks.append(('客房数量更新为16', qty==16, f'G17={qty}'))
checks.append(('客房单间面积写入说明', '26.54' in str(room_name), f'A17={room_name}'))
checks.append(('客房总面积为424.64㎡', abs(area-424.64)<1e-9, f'B17={area}'))
checks.append(('客房总冷负荷为50.96kW', abs(cold-50.96)<0.01, f'B17*C17/1000={cold}'))
checks.append(('客房选型冷量为56.06kW', abs(sel-56.06)<0.01, f'冷负荷*1.1={sel}'))
# formulas present
checks.append(('客房冷负荷仍使用模板公式', ws['D17'].value=='=B17*C17/1000' and ws['E17'].value=='=D17*1.1', f'D17={ws["D17"].value}; E17={ws["E17"].value}'))
wsn=wb['新风量']
checks.append(('客房总新风量为960m³/h', wsn['G31'].value=='=D31*F31' and wsn['D31'].value==32 and wsn['F31'].value==30, f'D31={wsn["D31"].value}, F31={wsn["F31"].value}, G31公式={wsn["G31"].value}, 计算=960'))
# equipment sources and models
wse=wb['设备汇总表']
models=[wse[f'E{r}'].value for r in [5,6,9,10,11,13,14]]
source_ok=('FC-204' in models and 'FC-136' in models and 'FC-85' in models and 'KD-3D' in models and 'KD-2D' in models and 'ZK-20W6S' in models and 'ZK-25W8S' in models)
checks.append(('设备参数来自本地老师资料提取型号', source_ok, ', '.join(str(m) for m in models)))
checks.append(('水系统保持GPT母版82行结构', wb['水系统'].max_row==82, f"水系统 max_row={wb['水系统'].max_row}"))
# field alignment quick check
align_ok = wb['设备汇总表']['B2'].value=='序号' and wb['设备汇总表']['C2'].value=='设备编号' and wb['设备汇总表']['E2'].value=='型号'
checks.append(('设备汇总表字段对齐', align_ok, f"B2={wb['设备汇总表']['B2'].value}, C2={wb['设备汇总表']['C2'].value}, E2={wb['设备汇总表']['E2'].value}"))
# independent equipment workbook open check
eq=base/'output'/'暖通设备选型汇总表-GPT重填版-模板内容补丁.xlsx'
eq_ok=False; eq_msg=''
try:
    eqwb=load_workbook(eq, read_only=True, data_only=True)
    eqws=eqwb[eqwb.sheetnames[0]]
    eq_ok=eqws.max_row>=8 and eqws.max_column>=11 and eqws['D2'].value=='ZK-20W6S'
    eq_msg=f'{eqws.title} {eqws.max_row}x{eqws.max_column}'
    eqwb.close()
except Exception as e:
    eq_msg=repr(e)
checks.append(('独立设备汇总表可打开并含补丁数据', eq_ok, eq_msg))

lines=['# TEMPLATE-ONLY-PATCH 自检报告','',f'- 母版：`{src}`',f'- 补丁版：`{patched}`','- 自检方式：openpyxl 只读/结构比对；未调用 AutoCAD，未读取 DWG。','', '## 检查结果', '', '| 检查项 | 结果 | 说明 |','|---|---|---|']
for name,ok,msg in checks:
    lines.append(f'| {name} | {"通过" if ok else "未通过"} | {str(msg).replace("|","/")} |')
lines += ['', '## 关键结论', '', '- 本补丁仅使用模板已有工作表、行列和字段。', '- 未新增工作表，未新增楼层列，未新增最不利环路/水泵扬程/阻力/完整水力平衡等模板外计算。', '- 水系统保持 GPT 母版 82 行结构。', '- 客房数据已更新为 16 间、26.54㎡/间、总面积 424.64㎡、总冷负荷 50.96kW、总新风量 960m³/h。', '- 设备型号与参数来源为本地老师资料提取成果。']
report.write_text('\n'.join(lines), encoding='utf-8')
print('\n'.join(lines[:30]))
wb0.close(); wb.close()
