# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook
p=Path('output/暖通课程设计计算表-GPT重填版-模板内容补丁.xlsx')
wb=load_workbook(p, data_only=False)
ws=wb['新风量']
ws['E29']=11/21.65
ws['D29']='=C29*E29'
ws['G29']='=D29*F29'
ws=wb['风-全空气']
ws['D54']='=F49'
try:
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = 'auto'
except Exception:
    pass
wb.save(p); wb.close()
print('restored formulas')
