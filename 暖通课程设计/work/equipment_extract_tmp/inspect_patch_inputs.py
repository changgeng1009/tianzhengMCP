# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook
base=Path.cwd()
patterns=[
 ('current', 'output/*GPT重填版.xlsx'),
 ('template', 'work/sandbox/*原始模板备份.xlsx'),
 ('teacher', 'REDACTED_CLASS 课程设计资料/**/REDACTED_CLASS-暖通课程设计计算表*.xlsx'),
 ('missing', 'output/*缺失数据*GPT重填版.md'),
 ('fcu','work/fcu_catalog_extracted.xlsx'),
 ('fresh','work/fresh_air_unit_catalog_extracted.xlsx'),
 ('zk1','work/zk_ahu_data.md'),
 ('zk2','REDACTED_CLASS 课程设计资料/**/workzk_ahu_data.md'),
 ('review','audit/gpt_refill_version_review.md'),
 ('risks','audit/gpt_refill_remaining_risks.md')]
found={}
for key,pat in patterns:
    matches=list(base.glob(pat))
    found[key]=matches[0] if matches else None
    print(key, found[key], found[key].exists() if found[key] else False, found[key].stat().st_size if found[key] else None)
for key in ['current','template']:
    p=found[key]
    if not p: continue
    wb=load_workbook(p, read_only=False, data_only=False)
    print('\nFILE',key,p)
    print('sheets', wb.sheetnames)
    for ws in wb.worksheets:
        print(ws.title, ws.max_row, ws.max_column, 'merged', len(ws.merged_cells.ranges))
    wb.close()
for key in ['missing','review','risks']:
    p=found[key]
    if p:
        print('\n---',key,p,'---')
        print(p.read_text(encoding='utf-8')[:1600])
