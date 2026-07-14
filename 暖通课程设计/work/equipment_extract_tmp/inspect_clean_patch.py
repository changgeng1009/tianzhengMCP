# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook
base=Path.cwd()
files=[base/'output'/'暖通课程设计计算表-GPT重填版-模板内容补丁.xlsx', base/'output'/'缺失数据与人工确认项-模板内容补丁版.md', base/'audit'/'template_only_patch_self_check.md']
for p in files:
    print(p, p.exists(), p.stat().st_size if p.exists() else None)
terms=['预留','需人工确认','待人工确认','按平面布置','按平面布置估算','待图纸布置后复核','待施工图复核','人工确认']
wb=load_workbook(files[0], data_only=False)
print('sheets', wb.sheetnames)
for ws in wb.worksheets:
    hits=[]
    for row in ws.iter_rows():
        for c in row:
            v=c.value
            if isinstance(v,str) and any(t in v for t in terms):
                hits.append((c.coordinate,v))
    if hits:
        print('\nSHEET', ws.title, 'hits', len(hits))
        for coord,v in hits[:80]:
            print(coord, repr(v))
print('\n室内外设计参数 rows 15-24')
ws=wb['室内外设计参数']
for r in range(15,25):
    print(r, [ws.cell(r,c).value for c in range(2,9)])
print('\n水系统 length cols C/K rows with text')
ws=wb['水系统']
for r in range(1, ws.max_row+1):
    for col in [3,11]:
        v=ws.cell(r,col).value
        if isinstance(v,str):
            print(ws.cell(r,col).coordinate, repr(v))
wb.close()
