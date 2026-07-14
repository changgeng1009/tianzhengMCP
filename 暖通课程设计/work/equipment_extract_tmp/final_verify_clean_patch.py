# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook
files=[Path('output')/'暖通课程设计计算表-GPT重填版-清洁补丁.xlsx', Path('output')/'Excel备注与人工确认项.md', Path('audit')/'excel_clean_patch_self_check.md']
for p in files:
    print(p, p.exists(), p.stat().st_size if p.exists() else None)
wb=load_workbook(files[0], read_only=True, data_only=False)
print('sheets', wb.sheetnames)
print('室内参数', [[wb['室内外设计参数'].cell(r,c).value for c in range(2,9)] for r in range(18,22)])
wb.close()
