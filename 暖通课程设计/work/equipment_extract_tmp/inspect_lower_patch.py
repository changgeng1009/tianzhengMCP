# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from pathlib import Path
p=Path('output/暖通课程设计计算表-GPT重填版.xlsx')
wb=load_workbook(p, data_only=False)
for sheet in ['风-FCU+新风','水系统','风-全空气']:
    ws=wb[sheet]
    print('\nSHEET',sheet, ws.max_row, ws.max_column)
    start=31 if sheet!='风-全空气' else 31
    for r in range(start, ws.max_row+1):
        vals=[]
        for c in range(1, ws.max_column+1):
            v=ws.cell(r,c).value
            if v is not None:
                vals.append(f'{ws.cell(r,c).coordinate}={str(v)[:80]}')
        if vals: print(' | '.join(vals))
wb.close()
