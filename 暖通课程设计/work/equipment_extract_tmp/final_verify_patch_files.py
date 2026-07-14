# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook
files=[
 Path('output')/'暖通课程设计计算表-GPT重填版-模板内容补丁.xlsx',
 Path('output')/'暖通设备选型汇总表-GPT重填版-模板内容补丁.xlsx',
 Path('output')/'Excel模板内容补丁说明.md',
 Path('output')/'缺失数据与人工确认项-模板内容补丁版.md',
 Path('audit')/'template_only_patch_self_check.md']
for p in files:
    print(p, p.exists(), p.stat().st_size if p.exists() else None)
for p in files[:2]:
    wb=load_workbook(p, read_only=True, data_only=False)
    print('xlsx', p.name, wb.sheetnames)
    wb.close()
print(Path('audit/template_only_patch_self_check.md').read_text(encoding='utf-8')[:1200])
