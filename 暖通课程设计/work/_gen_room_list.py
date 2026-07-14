import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'room_list'

header_font = Font(bold=True, size=11)
header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)
warn_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')

headers = ['Floor', 'RoomID', 'Name', 'Type', 'AC', 'Area_m2', 'Load_W_m2', 'ExtWall', 'Status']
for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.border = thin_border
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

rooms = [
    ['2F', '2F-CT01', 'Left CT', 'Dining Hall', 'All-Air', '', '380~480', '', 'TBD'],
    ['2F', '2F-CT02', 'Right CT', 'Dining Hall', 'All-Air', '', '380~480', '', 'TBD'],
    ['2F', '2F-BX01', 'Lux Box 1', 'Lux Box', 'FCU+OA', '', '300~360', '', 'TBD'],
    ['2F', '2F-BX02', 'Lux Box 2', 'Lux Box', 'FCU+OA', '', '300~360', '', 'TBD'],
    ['2F', '2F-BX03', 'Std Box 1', 'Std Box', 'FCU+OA', '', '300~360', '', 'TBD'],
    ['2F', '2F-BX04', 'Std Box 2', 'Std Box', 'FCU+OA', '', '300~360', '', 'TBD'],
    ['2F', '2F-BX05', 'Std Box 3', 'Std Box', 'FCU+OA', '', '300~360', '', 'TBD'],
    ['2F', '2F-BX06', 'Std Box 4', 'Std Box', 'FCU+OA', '', '300~360', '', 'TBD'],
    ['2F', '2F-BX07', 'Std Box 5', 'Std Box', 'FCU+OA', '', '300~360', '', 'TBD'],
    ['2F', '2F-BX08', 'Std Box 6', 'Std Box', 'FCU+OA', '', '300~360', '', 'TBD'],
    ['4F', '4F-K01', 'Guest Rm 01', 'Hotel Room', 'FCU+OA', '', 'TBD', '', 'TBD'],
    ['4F', '4F-K02', 'Guest Rm 02', 'Hotel Room', 'FCU+OA', '', 'TBD', '', 'TBD'],
    ['4F', '4F-K03', 'Guest Rm 03', 'Hotel Room', 'FCU+OA', '', 'TBD', '', 'TBD'],
    ['4F', '4F-K04', 'Guest Rm 04', 'Hotel Room', 'FCU+OA', '', 'TBD', '', 'TBD'],
    ['4F', '4F-K05', 'Guest Rm 05', 'Hotel Room', 'FCU+OA', '', 'TBD', '', 'TBD'],
    ['4F', '4F-K06', 'Guest Rm 06', 'Hotel Room', 'FCU+OA', '', 'TBD', '', 'TBD'],
    ['4F', '4F-K07', 'Guest Rm 07', 'Hotel Room', 'FCU+OA', '', 'TBD', '', 'TBD'],
    ['4F', '4F-K08', 'Guest Rm 08', 'Hotel Room', 'FCU+OA', '', 'TBD', '', 'TBD'],
    ['4F', '4F-K09', 'Guest Rm 09', 'Hotel Room', 'FCU+OA', '', 'TBD', '', 'TBD'],
    ['4F', '4F-K10', 'Guest Rm 10', 'Hotel Room', 'FCU+OA', '', 'TBD', '', 'TBD'],
]

for row_idx, room in enumerate(rooms, 2):
    for col_idx, val in enumerate(room, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center', vertical='center')

widths = [5, 12, 14, 14, 10, 10, 12, 8, 8]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

out = r'D:\CodexWork\天正\暖通课程设计\work\room_list_draft.xlsx'
wb.save(out)
print('Saved to ' + out)
