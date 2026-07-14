import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = '房间面积提取'

warn_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
ok_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                      top=Side(style='thin'), bottom=Side(style='thin'))

headers = ['楼层','编号','房间名','长(m)','宽(m)','面积(m2)','层高(m)','体积(m3)','外墙','数据来源','确认状态']
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal='center')

rooms = [
    ['2F','2F-CT01','中餐厅(左)','','','','','','','DWG 量取','待用户输入'],
    ['2F','2F-CT02','中餐厅(右)','','','','','','','DWG 量取','待用户输入'],
    ['2F','2F-BX01','豪华包厢1','','','','','','','DWG 量取','待用户输入'],
    ['2F','2F-BX02','豪华包厢2','','','','','','','DWG 量取','待用户输入'],
    ['2F','2F-BX03','普通包厢1','','','','','','','DWG 量取','待用户输入'],
    ['2F','2F-BX04','普通包厢2','','','','','','','DWG 量取','待用户输入'],
    ['2F','2F-BX05','普通包厢3','','','','','','','DWG 量取','待用户输入'],
    ['2F','2F-BX06','普通包厢4','','','','','','','DWG 量取','待用户输入'],
    ['2F','2F-BX07','普通包厢5','','','','','','','DWG 量取','待用户输入'],
    ['2F','2F-BX08','普通包厢6','','','','','','','DWG 量取','待用户输入'],
    ['4F','4F-K01','客房01','','','','','','','DWG 量取','待用户输入'],
    ['4F','4F-K02','客房02','','','','','','','DWG 量取','待用户输入'],
    ['4F','4F-K03','客房03','','','','','','','DWG 量取','待用户输入'],
    ['4F','4F-K04','客房04','','','','','','','DWG 量取','待用户输入'],
    ['4F','4F-K05','客房05','','','','','','','DWG 量取','待用户输入'],
    ['4F','4F-K06','客房06','','','','','','','DWG 量取','待用户输入'],
    ['4F','4F-K07','客房07','','','','','','','DWG 量取','待用户输入'],
    ['4F','4F-K08','客房08','','','','','','','DWG 量取','待用户输入'],
    ['4F','4F-K09','客房09','','','','','','','DWG 量取','待用户输入'],
    ['4F','4F-K10','客房10','','','','','','','DWG 量取','待用户输入'],
]

for r, room in enumerate(rooms, 2):
    for c, v in enumerate(room, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center')
        if c in [4,5,6,7,8]:
            cell.fill = warn_fill
        if v == '待用户输入':
            cell.fill = warn_fill

for i in range(1, 12):
    ws.column_dimensions[get_column_letter(i)].width = 14

# Instructions sheet
ws2 = wb.create_sheet('填写说明')
ws2['A1'] = '房间面积提取表 - 使用说明'
ws2['A1'].font = Font(bold=True, size=14)
notes = [
    '',
    '请在黄色标记的单元格中填入从 AutoCAD/天正 图纸中量取的数值。',
    '',
    '操作步骤：',
    '1. 用 AutoCAD 打开 D:\CodexWork\天正\暖通课程设计\REDACTED_CLASS 课程设计资料\REDACTED_CLASS 课程设计资料\图纸\需要设计的两层楼的酒店建筑图.dwg',
    '2. 仅需量取以下区域：二层中餐厅(2个)、豪华包厢(2个)、普通包厢(N个)、四层所有客房',
    '3. 用 AREA 命令或 DI 命令量取房间面积/长宽',
    '4. 将数值填入上表对应单元格',
    '5. 保存后告知 Codex 继续计算',
    '',
    '注意：本表为非自动化提取的备用方案。请勿修改原始 DWG。',
]
for i, note in enumerate(notes, 2):
    ws2[f'A{i}'] = note

wb.save(r'D:\CodexWork\天正\暖通课程设计\work\room_area_extraction.xlsx')
print('Done: room_area_extraction.xlsx')
